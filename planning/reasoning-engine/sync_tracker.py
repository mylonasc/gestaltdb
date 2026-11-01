"""Generate the Excel status view from EPIC_REASONING_ENGINE.md; no GitHub writes.

Only workbook I/O needs optional openpyxl. Parsing/validation is stdlib-only so
the roadmap gate and unit tests can check the authoritative Markdown cheaply.
"""

from collections import Counter
from dataclasses import dataclass
import argparse
from pathlib import Path
import re


ROOT = Path(__file__).resolve().parents[2]
STATUSES = ("Available", "Partial", "Experimental", "Deferred")


@dataclass(frozen=True)
class Review:
    timestamp: str
    revision: str
    published: str
    features: tuple[tuple[str, ...], ...]
    evidence: dict[str, str]
    roadmap: tuple[tuple[str | None, ...], ...]


def table(text, heading):
    section = text.split(heading + "\n", 1)[1].split("\n## ", 1)[0]
    rows = []
    for line in section.splitlines():
        if line.startswith("| "):
            if not line.startswith("| ---"):
                rows.append(tuple(cell.strip() for cell in line.strip("|").split("|")))
        elif rows:
            break
    return rows


def load_review(root=ROOT):
    text = (root / "EPIC_REASONING_ENGINE.md").read_text(encoding="utf-8")
    timestamp = re.search(r"\*\*Reviewed at:\*\* (\S+)", text).group(1)
    revision = re.search(r"\*\*Reviewed source:\*\* \[`([0-9a-f]{40})`\]", text).group(1)
    published = text.split("**Published baseline:** ", 1)[1].splitlines()[0].strip()
    published = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", published).replace("`", "")
    features = tuple(table(text, "## 7. Live 136-feature register (Excel source)")[1:])
    assert [row[0] for row in features] == [f"GR-{n:03}" for n in range(1, 137)], "GR IDs must be unique and ordered"
    evidence = {}
    for key, implementation, tests in table(text, "## 6. Evidence catalogue for the live register")[1:]:
        assert key not in evidence, f"Duplicate evidence key: {key}"
        evidence[key] = implementation.replace("`", "") + "; " + tests.replace("`", "")
        for path in re.findall(r"`([^`]+)`", implementation + tests):
            assert (root / path).is_file(), f"Missing evidence: {path}"
    for row in features:
        assert len(row) == 5 and row[2] in STATUSES, row
        assert row[4] and all(key in evidence for key in row[3].split(",")), row
    ledger = table(text, "### RE delivery ledger (replaces “implementation started” diaries)")[1:]
    assert [row[0] for row in ledger] == [f"RE-{n:02}" for n in range(1, 28)], "RE delivery IDs must be unique and ordered"
    index = table(text, "## 5. Completed feature issue index and original dependencies")[1:]
    dependencies = {row[0]: row[3] for row in index}
    roadmap = []
    for feature, merged, pr_commit, outcome in ledger:
        roadmap.append((feature, "Closed / delivered scoped outcome", outcome,
                        merged, pr_commit.replace("`", ""), dependencies[feature]))
    pending = table(text, "## 4. Pending work and deliberate deferrals")[1:]
    for issue, state, remaining in pending:
        number = re.search(r"#\d+", issue).group()
        roadmap.append((number, state, remaining, "Reviewed " + timestamp, None, "See remaining scope"))
    return Review(timestamp, revision, published, features, evidence, tuple(roadmap))


def feature_values(review, original_rows, owners):
    output = []
    for feature, construct, status, keys, boundary in review.features:
        old = original_rows[feature]
        assert construct == old[2], f"Original construct changed: {feature}"
        paths = " / ".join(review.evidence[key] for key in keys.split(","))
        notes = boundary + " Evidence: " + paths
        # Original acceptance/priority/layer remain intact, not weakened to fit delivery.
        output.append((feature, old[1], construct, status, old[4], old[5], old[6],
                       "https://github.com/mylonasc/gestaltdb/blob/main/EPIC_REASONING_ENGINE.md#7-live-136-feature-register-excel-source",
                       notes, owners.get(feature), "See completed RE ledger in epic", old[3],
                       review.timestamp, review.revision))
    return output


def dashboard_values(review):
    counts = Counter(row[2] for row in review.features)
    return [
        ("GESTALTDB · REASONING STATUS", "Generated from EPIC_REASONING_ENGINE.md"),
        ("Reviewed UTC", review.timestamp),
        ("Reviewed main source", review.revision),
        ("Published baseline", review.published),
        ("Status", "Count"),
        *((status, counts[status]) for status in STATUSES),
        ("Total constructs", len(review.features)),
        ("Meaning: Available", "Implemented within the row's declared scope; not native/performance certification."),
        ("Meaning: Partial", "Supporting implementation exists; original broader acceptance remains unmet."),
        ("Meaning: Experimental", "Explicit prototype only, not a general production capability."),
        ("Meaning: Deferred", "Not implemented/selected; not automatically scheduled or actively in progress."),
        ("Release boundary", "See epic section 3: published tag and post-tag changes are separate."),
        ("Roadmap disposition", "See epic sections 1 and 3; closed slices include approved defer/design outcomes."),
        ("Pending", "; ".join(f"{row[0]}: {row[1]}" for row in review.roadmap if row[0].startswith("#"))),
        ("Validation", "See epic section 8 for exact source command, versions, results and limitations."),
        ("Editing policy", "Edit the epic, then run sync_tracker.py. Save/close LibreOffice before writing."),
        ("Historical inputs", "Original audit and Original milestones preserve the previous workbook rows; git preserves its bytes."),
    ]


def _rows(sheet):
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def check_workbook(workbook, review):
    original = {row[0]: row for row in _rows(workbook["Original audit"])[1:]}
    current = _rows(workbook["Feature Register"])
    owners = {row[0]: row[9] for row in current[1:]}
    assert current[1:] == feature_values(review, original, owners), "Workbook feature/status/evidence drift"
    assert _rows(workbook["Dashboard"]) == dashboard_values(review), "Workbook dashboard drift"
    assert _rows(workbook["Roadmap"])[1:] == list(review.roadmap), "Workbook delivery/pending roadmap drift"


def _fill(sheet, rows, widths):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="203864")
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def main():
    import openpyxl

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Validate without writing")
    args = parser.parse_args()
    review = load_review()
    path = ROOT / "gestaltdb_reasoning_feature_tracker.xlsx"
    workbook = openpyxl.load_workbook(path)
    if args.check:
        check_workbook(workbook, review)
        print("Workbook matches all 136 epic rows, dashboard, review identity and delivery/pending ledger.")
        return
    lock = ROOT / (".~lock." + path.name + "#")
    if lock.exists():
        parser.error("Workbook is open in LibreOffice; save and close it before updating.")
    if "Original audit" not in workbook:
        workbook.copy_worksheet(workbook["Feature Register"]).title = "Original audit"
        workbook.copy_worksheet(workbook["Roadmap"]).title = "Original milestones"
    original = {row[0]: row for row in _rows(workbook["Original audit"])[1:]}
    owners = {row[0]: row[9] for row in _rows(workbook["Feature Register"])[1:]}
    headers = ("ID", "Workstream", "Construct / execution primitive", "Reviewed status", "Priority",
               "Original suggested layer", "Original acceptance criterion", "Status authority",
               "Verification / remaining boundary and evidence", "Owner", "Milestone reference",
               "Original audit status", "Review UTC", "Reviewed source commit")
    for name, rows, widths in (
        ("Feature Register", [headers, *feature_values(review, original, owners)],
         [12, 25, 38, 17, 10, 23, 65, 65, 110, 20, 35, 20, 25, 45]),
        ("Dashboard", dashboard_values(review), [28, 115]),
        ("Roadmap", [("ID", "State", "Delivered / remaining scope", "Merged UTC / review", "PR / commit", "Original prerequisites"),
                     *review.roadmap], [15, 38, 100, 30, 30, 35]),
    ):
        position = workbook.index(workbook[name])
        workbook.remove(workbook[name])
        sheet = workbook.create_sheet(name, position)
        _fill(sheet, rows, widths)
    colors = {"Available": "E2F0D9", "Partial": "FFF2CC", "Experimental": "DDEBF7", "Deferred": "E7E6E6"}
    from openpyxl.styles import PatternFill
    for row in workbook["Feature Register"].iter_rows(min_row=2):
        row[3].fill = PatternFill("solid", fgColor=colors[row[3].value])
        row[7].hyperlink = row[7].value
    workbook.active = workbook.index(workbook["Dashboard"])
    workbook.properties.description = "Generated status view; authority: EPIC_REASONING_ENGINE.md"
    check_workbook(workbook, review)
    # Stage and replace so an interrupted zip write cannot corrupt the original.
    import os
    import tempfile
    fd, temporary = tempfile.mkstemp(prefix=".reasoning-tracker-", suffix=".xlsx", dir=ROOT)
    os.close(fd)
    try:
        workbook.save(temporary)
        if lock.exists():
            raise RuntimeError("Workbook was opened in LibreOffice during generation; close it and retry.")
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    print("Updated 136 features, dashboard and delivery/pending ledger; original audit/acceptance preserved.")


if __name__ == "__main__":
    main()
