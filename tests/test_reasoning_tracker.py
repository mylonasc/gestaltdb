"""Keep the consolidated epic, generated workbook and tracking expectations aligned."""

from collections import Counter
from copy import deepcopy
import importlib
import hashlib
import json
from pathlib import Path
import shutil
import sys

import pytest


ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def tracker(monkeypatch):
    monkeypatch.syspath_prepend(str(ROOT / "planning/reasoning-engine"))
    return importlib.import_module("sync_tracker")


def test_register_has_all_original_ids_and_verified_evidence(tracker):
    review = tracker.load_review()
    assert [row[0] for row in review.features] == [f"GR-{n:03}" for n in range(1, 137)]
    assert set(row[2] for row in review.features) == set(tracker.STATUSES)
    assert len(review.roadmap) == 34
    assert review.roadmap[11][3] == "2026-10-21T06:47:24Z"
    assert review.roadmap[-1][0] == "#409"


def test_table_parser_does_not_swallow_later_diaries(tracker):
    text = "## Ledger\n\n| ID | Outcome |\n| --- | --- |\n| RE-01 | Delivered |\n\n### Later\n\n| Wrong | Table |\n"
    assert tracker.table(text, "## Ledger") == [("ID", "Outcome"), ("RE-01", "Delivered")]


@pytest.mark.parametrize("old,new", [
    ("| GR-002 |", "| GR-001 |"),
    ("| GR-001 | Stable entity & relationship IDs | Available |", "| GR-001 | Stable entity & relationship IDs | Complete |"),
    ("| Available | F,H |", "| Available | UNKNOWN |"),
    ("`tests/test_reasoning_arrow.py`", "`tests/missing_evidence.py`"),
])
def test_markdown_identity_status_and_evidence_errors_fail(tracker, monkeypatch, old, new):
    original_read = Path.read_text
    epic = ROOT / "EPIC_REASONING_ENGINE.md"
    text = epic.read_text()
    assert old in text

    def changed_read(path, *args, **kwargs):
        return text.replace(old, new, 1) if path == epic else original_read(path, *args, **kwargs)

    monkeypatch.setattr(Path, "read_text", changed_read)
    with pytest.raises(AssertionError):
        tracker.load_review()


def test_live_tracking_expects_closed_parent_and_release_child(tracker):
    tracking = importlib.import_module("check_tracking")
    snapshot = tracking.load()
    tracking.check_offline(snapshot)
    states = tracking.expected_states(snapshot)
    assert states["156"] == "closed"
    assert states["385"] == "closed"
    assert states["194"] == "open"
    assert all(states[str(n)] == "closed" for n in range(391, 397))
    assert 385 in snapshot["expected_sub_issues"]["156"]


def test_tracking_rejects_stale_review_identity(tracker):
    tracking = importlib.import_module("check_tracking")
    snapshot = deepcopy(tracking.load())
    snapshot["reviewed_at"] = "2026-10-25T00:00:00Z"
    with pytest.raises(AssertionError):
        tracking.check_offline(snapshot)


def test_dashboard_counts_and_pending_come_from_epic(tracker):
    review = tracker.load_review()
    dashboard = dict(tracker.dashboard_values(review))
    assert dashboard["Total constructs"] == 136
    counts = Counter(row[2] for row in review.features)
    assert all(dashboard[status] == counts[status] for status in tracker.STATUSES)
    assert "#194" in dashboard["Pending"]
    assert review.published in dashboard.values()


def test_original_audit_digest_is_self_contained_without_git(tracker):
    checker = importlib.import_module("check_roadmap")
    plan = json.loads((ROOT / "planning/reasoning-engine/roadmap.json").read_text())
    original = checker.workbook_rows(ROOT / "gestaltdb_reasoning_feature_tracker.xlsx", sheet_name="Original audit")
    digest = hashlib.sha256(json.dumps(original, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()
    assert digest == plan["historical_workbook"]["original_audit_rows_sha256"]
    live = checker.workbook_rows(ROOT / "gestaltdb_reasoning_feature_tracker.xlsx")
    for row, baseline, expected in zip(live[1:], original[1:], tracker.load_review().features):
        assert row[3] == expected[2]
        assert row[6] == baseline[6]
        assert row[11] == baseline[3]


@pytest.fixture
def workbook():
    openpyxl = pytest.importorskip("openpyxl", reason="optional Excel synchronization checks")
    return openpyxl.load_workbook(ROOT / "gestaltdb_reasoning_feature_tracker.xlsx")


def test_generated_workbook_and_historical_criteria(tracker, workbook):
    review = tracker.load_review()
    tracker.check_workbook(workbook, review)
    original = list(workbook["Original audit"].values)[1:]
    current = list(workbook["Feature Register"].values)[1:]
    for previous, row in zip(original, current):
        assert [row[n] for n in (0, 1, 2, 4, 5, 6)] == [previous[n] for n in (0, 1, 2, 4, 5, 6)]
        assert row[11] == previous[3]


@pytest.mark.parametrize("sheet,cell,value", [
    ("Feature Register", "D2", "Deferred"),
    ("Feature Register", "G2", "Weakened acceptance"),
    ("Feature Register", "I2", "No evidence"),
    ("Feature Register", "N2", "stale-source"),
    ("Dashboard", "B6", -1),
    ("Roadmap", "D2", "wrong-merge-date"),
])
def test_workbook_drift_is_detected(tracker, workbook, sheet, cell, value):
    workbook[sheet][cell] = value
    with pytest.raises(AssertionError):
        tracker.check_workbook(workbook, tracker.load_review())


def test_excel_regeneration_roundtrip_preserves_owners(tracker, workbook, tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    import openpyxl

    path = tmp_path / "gestaltdb_reasoning_feature_tracker.xlsx"
    workbook["Feature Register"]["J2"] = "example-owner"
    workbook.save(path)
    monkeypatch.setattr(tracker, "ROOT", tmp_path)
    monkeypatch.setattr(sys, "argv", ["sync_tracker.py"])
    tracker.main()
    updated = openpyxl.load_workbook(path)
    tracker.check_workbook(updated, tracker.load_review())
    assert updated["Feature Register"]["J2"].value == "example-owner"
    assert updated.sheetnames.count("Original audit") == 1


def test_open_workbook_is_never_overwritten(tracker, tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    path = tmp_path / "gestaltdb_reasoning_feature_tracker.xlsx"
    shutil.copyfile(ROOT / path.name, path)
    before = path.read_bytes()
    (tmp_path / (".~lock." + path.name + "#")).write_text("user-owned lock")
    monkeypatch.setattr(tracker, "ROOT", tmp_path)
    monkeypatch.setattr(sys, "argv", ["sync_tracker.py"])
    with pytest.raises(SystemExit):
        tracker.main()
    assert path.read_bytes() == before
